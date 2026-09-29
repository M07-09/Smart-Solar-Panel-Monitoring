"""Build the Word version of the report from report_content.py (same content as the PDF).

python-docx writes the document; Microsoft Word (through pywin32) then fills in the table of
contents and page numbers and saves it. Usage:  python build_docx.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from report_content import BLOCKS, META

HERE = Path(__file__).resolve().parent
OUT = HERE / "Smart_Solar_Monitoring_Report.docx"

NAVY, BLUE, AMBER, GRAY = RGBColor(0x11, 0x30, 0x5C), RGBColor(0x1F, 0x5F, 0xA8), "F0A202", RGBColor(0x5B, 0x64, 0x72)
CONTENT_W_CM, MAX_FIG_H_CM = 17.0, 21.0
INLINE = re.compile(r"\*\*(.+?)\*\*|(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])")


# ---------------------------------------------------------------- low-level helpers
def add_runs(par, text, size=None, color=None, italic=False):
    """Add text with **bold** and *italic* markup as runs."""
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            _run(par, text[pos:m.start()], size, color, italic=italic)
        _run(par, m.group(1) or m.group(2), size, color, bold=m.group(1) is not None, italic=italic or m.group(2) is not None)
        pos = m.end()
    if pos < len(text):
        _run(par, text[pos:], size, color, italic=italic)


def _run(par, text, size, color, bold=False, italic=False):
    r = par.add_run(text)
    r.bold, r.italic = bold or None, italic or None
    if size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color


def bottom_border(par, color=AMBER, size=12):
    ppr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(size)), ("w:space", "4"), ("w:color", color)):
        b.set(qn(k), v)
    bdr.append(b)
    ppr.append(bdr)


def add_field(par, code, size=8, color=RGBColor(0x8A, 0x8F, 0x98)):
    """Insert a Word field (PAGE, NUMPAGES, TOC ...)."""
    r = par.add_run()
    r.font.size, r.font.color.rgb = Pt(size), color
    for tag, attr in (("w:fldChar", "begin"), ("w:instrText", code), ("w:fldChar", "separate"),
                      ("w:t", "1"), ("w:fldChar", "end")):
        el = OxmlElement(tag)
        if tag == "w:fldChar":
            el.set(qn("w:fldCharType"), attr)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = attr
        r._r.append(el)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", fill)):
        shd.set(qn(k), v)
    tcpr.append(shd)


def row_flag(row, tag):
    el = OxmlElement(tag)
    row._tr.get_or_add_trPr().append(el)


def caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    add_runs(p, text, size=9, color=GRAY, italic=True)


# ---------------------------------------------------------------- blocks
def add_table(doc, header, rows, cap):
    lens = [max(len(str(h)), *(len(str(r[i]).replace("*", "")) for r in rows)) for i, h in enumerate(header)]
    weights = [min(max(n, 6), 60) for n in lens]
    widths = [max(CONTENT_W_CM * w / sum(weights), CONTENT_W_CM * 0.12) for w in weights]
    widths = [w * CONTENT_W_CM / sum(widths) for w in widths]

    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, values in enumerate([header] + rows):
        row = t.rows[ri]
        row_flag(row, "w:tblHeader" if ri == 0 else "w:cantSplit")
        for ci, v in enumerate(values):
            cell = row.cells[ci]
            cell.width = Cm(widths[ci])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if ri == 0:
                shade(cell, "11305C")
                _run(p, str(v), 9.5, RGBColor(0xFF, 0xFF, 0xFF), bold=True)
            else:
                if ri % 2 == 0:
                    shade(cell, "F5F7FB")
                add_runs(p, str(v), size=9.5)
    caption(doc, cap)


def add_figure(doc, rel, cap, pct):
    w, h = Image.open(HERE / rel).size
    width_cm = min(CONTENT_W_CM * pct / 100, MAX_FIG_H_CM * w / h)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(str(HERE / rel), width=Cm(width_cm))
    caption(doc, cap)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, side, Cm(2))

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.2
    for name, size, color, before in (("Heading 1", 16, NAVY, 18), ("Heading 2", 13, BLUE, 12)):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold, st.font.color.rgb = "Calibri", Pt(size), True, color
        rfonts = st.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)  # theme fonts (Cambria) would override Calibri
        st.paragraph_format.space_before, st.paragraph_format.space_after = Pt(before), Pt(6)
        st.paragraph_format.keep_with_next = True

    # footer: title  ....  Page X of Y
    fp = sec.footer.paragraphs[0]
    tabs = fp.paragraph_format.tab_stops
    for inherited in (Cm(8.26), Cm(16.51)):  # the Footer style's own centre/right tabs (Letter width)
        tabs.add_tab_stop(inherited, WD_TAB_ALIGNMENT.CLEAR)
    tabs.add_tab_stop(Cm(CONTENT_W_CM), WD_TAB_ALIGNMENT.RIGHT)
    _run(fp, META["title"] + "\tPage ", 8, RGBColor(0x8A, 0x8F, 0x98))
    add_field(fp, "PAGE")
    _run(fp, " of ", 8, RGBColor(0x8A, 0x8F, 0x98))
    add_field(fp, "NUMPAGES")

    # cover
    doc.add_paragraph().paragraph_format.space_before = Pt(150)
    _run(doc.add_paragraph(), META["subtitle"].upper(), 10, RGBColor(0xB3, 0x6B, 0x00), bold=True)
    p = doc.add_paragraph()
    _run(p, META["title"], 28, NAVY, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(40)
    _run(p, "Finding which solar inverter is losing energy, why, and what to do about it", 13,
         RGBColor(0x4A, 0x55, 0x68))
    for k, v in (("Student", META["student"]), ("Course", META["course"]),
                 ("Instructor", META["instructor"]), ("Date", META["date"])):
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(3.8))
        _run(p, k, 11, RGBColor(0x6B, 0x72, 0x80))
        _run(p, "\t" + v, 11, None)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(40)
    _run(p, "Random Forest · ResNet-18 (CNN) · Gemma 2 Transformer · FastAPI · Streamlit", 9.5,
         RGBColor(0x4A, 0x55, 0x68))
    p.add_run().add_break(WD_BREAK.PAGE)

    # contents (a styled paragraph, not a heading, so it does not list itself)
    p = doc.add_paragraph()
    _run(p, "Contents", 16, NAVY, bold=True)
    bottom_border(p)
    add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u', size=11, color=None)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    for b in BLOCKS:
        kind = b[0]
        if kind in ("h1", "h2"):
            h = doc.add_heading(b[1], level=1 if kind == "h1" else 2)
            if kind == "h1":
                bottom_border(h)
        elif kind == "p":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, b[1])
        elif kind in ("bullets", "numbered"):
            for i, item in enumerate(b[1], 1):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(0.9), Cm(-0.5)
                p.paragraph_format.space_after = Pt(4)
                _run(p, ("•" if kind == "bullets" else f"{i}.") + "\t", None, None)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
                add_runs(p, item)
        elif kind == "table":
            add_table(doc, b[1], b[2], b[3])
        elif kind == "figure":
            add_figure(doc, b[1], b[2], b[3])
        elif kind == "pagebreak":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    doc.core_properties.title, doc.core_properties.author = META["title"], META["student"]
    doc.core_properties.comments = ""  # python-docx writes its own name here by default
    doc.save(OUT)


def update_fields_with_word(preview_pdf: Path | None = None):
    """Let Microsoft Word fill in the contents page and page numbers, then save."""
    import win32com.client

    word = win32com.client.DispatchEx("Word.Application")
    word.Visible, word.DisplayAlerts = False, 0
    try:
        doc = word.Documents.Open(str(OUT))
        for toc in doc.TablesOfContents:
            toc.Update()
        doc.Fields.Update()
        doc.Save()
        if preview_pdf:
            doc.ExportAsFixedFormat(str(preview_pdf), 17)  # 17 = PDF
        doc.Close()
    except Exception as e:  # Word sometimes drops the COM link after exporting; the file is already saved
        print("Word:", e)
    finally:
        try:
            word.Quit()
        except Exception:
            pass


if __name__ == "__main__":
    import sys

    build()
    update_fields_with_word(Path(sys.argv[1]) if len(sys.argv) > 1 else None)
    print("saved", OUT)
